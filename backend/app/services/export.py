"""
报表导出服务
支持导出考核数据到 Excel
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from datetime import datetime
from sqlalchemy.orm import Session
from sqlalchemy import select

from app.models import (
    AssessmentRecord, AssessmentPlan, Employee, Department,
    EvalDetail, EvalTemplate, ApprovalLog
)


class ExportService:
    """报表导出服务"""

    @staticmethod
    def _apply_header_style(ws, row_num: int, col_count: int):
        """应用表头样式"""
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col in range(1, col_count + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

    @staticmethod
    def _apply_cell_style(ws, row_num: int, col_count: int):
        """应用单元格样式"""
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col in range(1, col_count + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='center')

    @staticmethod
    def _auto_adjust_column_width(ws):
        """自动调整列宽"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    @staticmethod
    def export_assessment_records(db: Session, plan_id: int = None) -> BytesIO:
        """
        导出考核记录
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "考核记录"

        # 查询数据
        stmt = (
            select(
                AssessmentRecord,
                Employee.name,
                Employee.wecom_userid,
                Department.name,
                Employee.position,
                AssessmentPlan.name,
                AssessmentPlan.cycle_type
            )
            .join(Employee, AssessmentRecord.employee_id == Employee.id)
            .outerjoin(Department, Employee.dept_id == Department.id)
            .join(AssessmentPlan, AssessmentRecord.plan_id == AssessmentPlan.id)
            .order_by(AssessmentRecord.created_at.desc())
        )

        if plan_id:
            stmt = stmt.where(AssessmentRecord.plan_id == plan_id)

        results = db.execute(stmt).all()

        # 表头
        headers = [
            "考核计划", "周期类型", "员工姓名", "企微ID", "部门", "职位",
            "当前状态", "当前步骤", "最终得分", "最终评语", "创建时间"
        ]
        ws.append(headers)
        ExportService._apply_header_style(ws, 1, len(headers))

        # 数据行
        for row_num, row in enumerate(results, start=2):
            record, emp_name, emp_userid, dept_name, position, plan_name, cycle_type = row
            ws.append([
                plan_name,
                cycle_type,
                emp_name,
                emp_userid,
                dept_name or "未分配",
                position or "未设置",
                record.status,
                record.current_step,
                record.final_score if record.final_score else "",
                record.final_comment if record.final_comment else "",
                record.created_at.strftime("%Y-%m-%d %H:%M:%S") if record.created_at else ""
            ])
            ExportService._apply_cell_style(ws, row_num, len(headers))

        ExportService._auto_adjust_column_width(ws)

        # 保存到内存
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_assessment_detail(db: Session, record_id: int) -> BytesIO:
        """
        导出单个考核详情（包含所有评分维度）
        """
        wb = Workbook()

        # 查询考核记录
        record = db.execute(
            select(AssessmentRecord).where(AssessmentRecord.id == record_id)
        ).scalar_one_or_none()

        if not record:
            raise ValueError("考核记录不存在")

        # 查询员工信息
        emp_row = db.execute(
            select(Employee, Department.name)
            .outerjoin(Department, Employee.dept_id == Department.id)
            .where(Employee.id == record.employee_id)
        ).one_or_none()

        employee = emp_row[0] if emp_row else None
        dept_name = emp_row[1] if emp_row else None

        # 查询计划和模板
        plan_row = db.execute(
            select(AssessmentPlan, EvalTemplate)
            .outerjoin(EvalTemplate, AssessmentPlan.template_id == EvalTemplate.id)
            .where(AssessmentPlan.id == record.plan_id)
        ).one_or_none()

        plan = plan_row[0] if plan_row else None
        template = plan_row[1] if plan_row else None

        # Sheet 1: 基本信息
        ws1 = wb.active
        ws1.title = "基本信息"
        ws1.append(["考核基本信息"])
        ws1.merge_cells('A1:B1')
        ws1['A1'].font = Font(bold=True, size=14)
        ws1['A1'].alignment = Alignment(horizontal='center')

        info_data = [
            ["员工姓名", employee.name if employee else ""],
            ["企微ID", employee.wecom_userid if employee else ""],
            ["部门", dept_name or ""],
            ["职位", employee.position if employee else ""],
            ["考核计划", plan.name if plan else ""],
            ["周期类型", plan.cycle_type if plan else ""],
            ["模板名称", template.name if template else ""],
            ["当前状态", record.status],
            ["当前步骤", record.current_step],
            ["最终得分", record.final_score if record.final_score else ""],
            ["最终评语", record.final_comment if record.final_comment else ""],
        ]

        for row in info_data:
            ws1.append(row)

        ExportService._auto_adjust_column_width(ws1)

        # Sheet 2: 评分详情
        ws2 = wb.create_sheet("评分详情")
        headers = ["维度名称", "维度权重", "维度类型", "评估人角色", "评估人", "评分", "评语"]
        ws2.append(headers)
        ExportService._apply_header_style(ws2, 1, len(headers))

        # 查询评分详情
        details = db.execute(
            select(EvalDetail, Employee.name)
            .join(Employee, EvalDetail.evaluator_id == Employee.id)
            .where(EvalDetail.record_id == record_id)
            .order_by(EvalDetail.created_at)
        ).all()

        for row_num, (detail, evaluator_name) in enumerate(details, start=2):
            ws2.append([
                detail.dimension_name,
                detail.dimension_weight,
                detail.dimension_type,
                detail.evaluator_role,
                evaluator_name,
                detail.score if detail.score else "",
                detail.comment if detail.comment else ""
            ])
            ExportService._apply_cell_style(ws2, row_num, len(headers))

        ExportService._auto_adjust_column_width(ws2)

        # Sheet 3: 审批流水
        ws3 = wb.create_sheet("审批流水")
        headers = ["来源步骤", "目标步骤", "操作", "审批人", "意见", "时间"]
        ws3.append(headers)
        ExportService._apply_header_style(ws3, 1, len(headers))

        # 查询审批日志
        logs = db.execute(
            select(ApprovalLog, Employee.name)
            .join(Employee, ApprovalLog.approver_id == Employee.id)
            .where(ApprovalLog.record_id == record_id)
            .order_by(ApprovalLog.created_at)
        ).all()

        for row_num, (log, approver_name) in enumerate(logs, start=2):
            ws3.append([
                log.from_step or "",
                log.to_step or "",
                log.action,
                approver_name,
                log.comment or "",
                log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else ""
            ])
            ExportService._apply_cell_style(ws3, row_num, len(headers))

        ExportService._auto_adjust_column_width(ws3)

        # 保存到内存
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_plan_summary(db: Session, plan_id: int) -> BytesIO:
        """
        导出考核计划汇总报表
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "考核汇总"

        # 查询计划信息
        plan = db.execute(
            select(AssessmentPlan).where(AssessmentPlan.id == plan_id)
        ).scalar_one_or_none()

        if not plan:
            raise ValueError("考核计划不存在")

        # 标题
        ws.append([f"考核计划汇总报表 - {plan.name}"])
        ws.merge_cells('A1:J1')
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([])  # 空行

        # 表头
        headers = [
            "序号", "员工姓名", "部门", "职位", "状态", "当前步骤",
            "自评分", "上级评分", "最终得分", "完成时间"
        ]
        ws.append(headers)
        ExportService._apply_header_style(ws, 3, len(headers))

        # 查询数据
        records = db.execute(
            select(
                AssessmentRecord,
                Employee.name,
                Department.name,
                Employee.position
            )
            .join(Employee, AssessmentRecord.employee_id == Employee.id)
            .outerjoin(Department, Employee.dept_id == Department.id)
            .where(AssessmentRecord.plan_id == plan_id)
            .order_by(Department.name, Employee.name)
        ).all()

        for idx, (record, emp_name, dept_name, position) in enumerate(records, start=1):
            # 查询自评分
            self_score = db.execute(
                select(EvalDetail.score)
                .where(
                    EvalDetail.record_id == record.id,
                    EvalDetail.evaluator_role == "self"
                )
            ).scalars().first()

            # 查询上级评分
            manager_score = db.execute(
                select(EvalDetail.score)
                .where(
                    EvalDetail.record_id == record.id,
                    EvalDetail.evaluator_role == "manager"
                )
            ).scalars().first()

            ws.append([
                idx,
                emp_name,
                dept_name or "未分配",
                position or "未设置",
                record.status,
                record.current_step,
                self_score if self_score else "",
                manager_score if manager_score else "",
                record.final_score if record.final_score else "",
                record.updated_at.strftime("%Y-%m-%d") if record.updated_at else ""
            ])
            ExportService._apply_cell_style(ws, idx + 3, len(headers))

        ExportService._auto_adjust_column_width(ws)

        # 保存到内存
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output


# 单例
export_service = ExportService()

